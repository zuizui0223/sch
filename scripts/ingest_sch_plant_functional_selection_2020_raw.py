"""Recover large original source meta databases for REAL macro-comparative analysis.

Caruso, Maherali & Martin (2020), Int J Plant Sci 181:44–55,
doi:10.1086/706199.
Original published Dryad dataset doi:10.5061/dryad.sn02v6x0q,
unchanged Zenodo archived copy record 5000217.

The experimental MASTER, observational, and intentionally DUPLICATED
workbooks have different source roles. Never concatenate them or pretend
that copied coefficients represent independent plant populations.
Scientific outcome tests do not run here: schema audit comes FIRST.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import zipfile
from pathlib import Path
from urllib.request import Request,urlopen
from xml.etree import ElementTree as ET

ORIGINAL_DOI="10.5061/dryad.sn02v6x0q"
ARTICLE_DOI="10.1086/706199"
ZENODO=5000217
FILENAMES={
  "Dryad_ReadMe.docx":("0e9a28b40290a3f07ba3c7d7dd2b5a3d",13_000),
  "Plant_functional_traits_experimental_database_duplicated_Dryad.xlsx":(
      "ce7f0c0112ee588dddea58218223645a",152_000),
  "Plant_functional_traits_experimental_database_master_Dryad.xlsx":(
      "7f680c344c0a418a3d8742649de24d52",83_500),
  "Plant_functional_traits_observational_database_Dryad.xlsx":(
      "248499857ad995490d8f9ac7dfbc94f8",124_000),
}
MAX_BYTES=500_000
UA="SCH-macro-original-plant-trait-selection/1.0"


def fetch(name:str)->tuple[bytes,str]:
    if name not in FILENAMES:raise ValueError("unregistered public source file")
    urls=(
        f"https://zenodo.org/api/records/{ZENODO}/files/{name}/content",
        f"https://zenodo.org/records/{ZENODO}/files/{name}?download=1",
    )
    errors=[]
    for url in urls:
        try:
            request=Request(url,headers={"User-Agent":UA,"Accept":"*/*"})
            with urlopen(request,timeout=40) as response:
                raw=response.read(MAX_BYTES+1)
            if len(raw)>MAX_BYTES:
                raise ValueError("file exceeds published reasonable size")
            d=hashlib.md5(raw).hexdigest()  # archival identity check only
            if d!=FILENAMES[name][0]:
                raise ValueError("file bytes do not match published archive MD5")
            if abs(len(raw)-FILENAMES[name][1])>20_000:
                raise ValueError("archive raw size unexpectedly changed")
            return raw,url
        except Exception as exc:
            errors.append(f"{type(exc).__name__} {exc}")
    raise ConnectionError(f"cannot fetch source-identical {name}: "+ "; ".join(errors))


def workbook_schema(path:Path)->list[dict]:
    import openpyxl
    book=openpyxl.load_workbook(path,read_only=True,data_only=True)
    report=[]
    for sh in book.worksheets:
        rows=sh.iter_rows(values_only=True)
        first=[list(next(rows,())) for i in range(min(8,sh.max_row or 0))]
        heading=[str(x) if x is not None else "" for x in first[0]] if first else []
        result={
            "sheet":sh.title,
            "n_rows_including_header":sh.max_row,
            "n_columns":sh.max_column,
            "source_column_names":heading[:150],
            "first_seven_source_rows":[
                [str(v)[:140] if v is not None else ""
                 for v in row[:min(35,len(row))]]
                for row in first[1:]
            ],
        }
        report.append(result)
    book.close()
    return report


def docx_notes(path:Path)->dict:
    ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(path) as z:
        root=ET.fromstring(z.read("word/document.xml"))
    paras=[]
    for p in root.findall(".//w:p",ns):
        t="".join(x.text or "" for x in p.findall(".//w:t",ns)).strip()
        if t:paras.append(t)
    return {
        "n_original_dictionary_paragraphs":len(paras),
        "original_dictionary_text_first_30000_chars":"\n".join(paras)[:30000],
    }


def main()->None:
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument("--outdir",type=Path,
                   default=Path("artifacts/plant_functional_selection_original"))
    a=p.parse_args();a.outdir.mkdir(parents=True,exist_ok=True)
    source={"data_doi":ORIGINAL_DOI,"paper_doi":ARTICLE_DOI,
            "original_archive_Zenodo_record":ZENODO,
            "n_original_source_files_expected":len(FILENAMES),
            "source_role":{
                "experimental_master":"UN-DUPLICATED ORIGINAL EXPERIMENTAL META DATABASE",
                "experimental_duplicated":"INTENTIONALLY DUPLICATED ORIGINAL AUTHOR ANALYSIS DATA",
                "observational":"OBSERVATIONAL POPULATION SPATIOTEMPORAL ESTIMATES",
            },
            "files":[],
            "status":"SOURCE_IDENTICAL_META_ORIGINAL_INGRESS_SCHEMA_ONLY",
            "new_beta_or_sign_analysis_executed":False}
    for name in FILENAMES:
        raw,url=fetch(name)
        dest=a.outdir/name;dest.write_bytes(raw)
        r={"name":name,"source_url":url,"source_md5":FILENAMES[name][0],
           "bytes":len(raw),"sha256":hashlib.sha256(raw).hexdigest()}
        if name.endswith(".xlsx"):r["workbook"]=workbook_schema(dest)
        if name.endswith(".docx"):r["variable_dictionary"]=docx_notes(dest)
        source["files"].append(r)
    out=a.outdir/"original_source_schema.json"
    out.write_text(json.dumps(source,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print("BEGIN_SCH_FUNCTIONAL_PLANT_META_SCHEMA")
    print(out.read_text(encoding="utf-8"))
    print("END_SCH_FUNCTIONAL_PLANT_META_SCHEMA")


if __name__=="__main__":
    main()

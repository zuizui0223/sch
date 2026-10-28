#!/usr/bin/env python3
"""Download & inspect original *meta-analysis compiled* source records.

Published source:
Caruso, Eisen, Martin & Sletvold (2018) meta-analytic database
doi:10.5061/dryad.2v8c5g0

The two files have distinct roles:
Exp_stud_NOTdup_Dryad.xls is the unduplicated original analytic source.
Exp_stud_dup_Dryad.xlsx contains INTENTIONAL DUPLICATES introduced for
the original authors' analyses; never combine these as distinct evidence.

Only source metadata, exact file hashes and schema are emitted here.
No new effects, contrasts or ecological prevalence estimators are computed.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import urllib.parse
import urllib.request
from pathlib import Path

DOI="10.5061/dryad.2v8c5g0"
EXPECTED_FILES={
  "Exp_stud_NOTdup_Dryad.xls":(300_000,600_000),
  "Exp_stud_dup_Dryad.xlsx":(150_000,350_000),
}
BASE="https://datadryad.org"
AGENT="SCH-floral-selection-meta-original-data-audit/1.0"
MAX_BYTES=2*1024*1024


def get_bytes(url:str, max_bytes:int=MAX_BYTES) -> bytes:
    request=urllib.request.Request(
        url,headers={"User-Agent":AGENT,"Accept":"application/json,*/*"})
    with urllib.request.urlopen(request,timeout=40) as f:
        b=f.read(max_bytes+1)
    if len(b)>max_bytes:
        raise ValueError("Unexpected giant source; no unbounded downloads")
    return b


def get_json(url:str)->dict:
    b=get_bytes(url)
    x=json.loads(b.decode("utf-8-sig"))
    if not isinstance(x,dict):
        raise ValueError("Dryad API response was not an object")
    return x


def absolute_api(relative:str)->str:
    if relative.startswith("https://datadryad.org/"):
        return relative
    if relative.startswith("/api/v2/"):
        return BASE+relative
    raise ValueError("unexpected external or non-API source link")


def get_dataset()->tuple[dict,dict,list[dict]]:
    encoded=urllib.parse.quote("doi:"+DOI,safe="")
    doc=get_json(f"{BASE}/api/v2/datasets/{encoded}")
    if doc.get("identifier","").removeprefix("doi:") != DOI:
        raise ValueError("unexpected original Dryad DOI identity")
    rel=doc.get("_links",{}).get("stash:version",{}).get("href")
    if not rel:
        versions=get_json(f"{BASE}/api/v2/datasets/{encoded}/versions")
        candidates=versions.get("_embedded",{}).get("stash:versions",[])
        if not candidates:
            raise ValueError("original dataset has no version link")
        selected=max(candidates,key=lambda x:x.get("versionNumber",0))
        rel=selected["_links"]["self"]["href"]
    version=get_json(absolute_api(rel))
    files_rel=version.get("_links",{}).get("stash:files",{}).get("href")
    if not files_rel:
        files_rel="/api/v2/versions/"+str(version["id"])+"/files"
    docfiles=get_json(absolute_api(files_rel))
    files=docfiles.get("_embedded",{}).get("stash:files",[])
    if set(f.get("path") for f in files)!=set(EXPECTED_FILES):
        raise ValueError("Original Dryad file list differs or unexpectedly duplicates: "+
                         str([f.get("path") for f in files]))
    return doc,version,files


def inspect_workbook(path:Path)->dict:
    name=path.name
    if name.endswith(".xls"):
        import xlrd
        book=xlrd.open_workbook(str(path),on_demand=True)
        source="xlrd"
        sheets=[]
        for sh in book.sheets():
            cells=[]
            for ri in range(min(sh.nrows,6)):
                cells.append([str(sh.cell_value(ri,j))[:150]
                              for j in range(min(sh.ncols,45))])
            sheets.append({
                "name":sh.name,"n_rows":sh.nrows,"n_columns":sh.ncols,
                "header_and_first_five_rows":cells,
            })
    elif name.endswith(".xlsx"):
        import openpyxl
        book=openpyxl.load_workbook(path,read_only=True,data_only=True)
        source="openpyxl"
        sheets=[]
        for sh in book.worksheets:
            first=[]
            for row in sh.iter_rows(min_row=1,max_row=6,values_only=True):
                first.append([str(v)[:150] if v is not None else ""
                              for v in list(row)[:45]])
            sheets.append({
                "name":sh.title,"n_rows":sh.max_row,"n_columns":sh.max_column,
                "header_and_first_five_rows":first,
            })
        book.close()
    else:
        raise ValueError("unknown original source file")
    return {"file":name,"reader":source,"sheets":sheets}


def main()->None:
    ap=argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--outdir",type=Path,default=Path("artifacts/sch_caruso_meta_raw"))
    args=ap.parse_args()
    args.outdir.mkdir(parents=True,exist_ok=True)
    dataset,version,files=get_dataset()
    result={
        "dataset_doi":DOI,
        "dataset_title":dataset.get("title"),
        "api_dataset_id":dataset.get("id"),
        "version_id":version.get("id"),
        "version_number":version.get("versionNumber"),
        "n_expected_files":len(EXPECTED_FILES),
        "source_file_roles":{
            "Exp_stud_NOTdup_Dryad.xls":"UNDUPLICATED_PRIMARY_META_COMPILATION",
            "Exp_stud_dup_Dryad.xlsx":"INTENTIONALLY_DUPLICATED_PUBLISHED_META_ANALYSIS_INPUT",
        },
        "file_receipts":[],
        "status":"ORIGINAL_PUBLISHED_META_DATABASE_INGESTION_ONLY",
        "no_new_statistical_analysis_performed":True,
    }
    for f in sorted(files,key=lambda f:f["path"]):
        filename=f["path"]
        rel=f.get("_links",{}).get("stash:download",{}).get("href")
        if not rel:
            raise ValueError("Dryad original file missing download link")
        url=absolute_api(rel)
        b=get_bytes(url)
        low,high=EXPECTED_FILES[filename]
        if not low<=len(b)<=high:
            raise ValueError("Unexpected original file size "+filename+" "+str(len(b)))
        algorithm=(f.get("digestType") or "").lower().replace("-","")
        expected=(f.get("digest") or "").lower()
        if algorithm in {"md5","sha256"}:
            actual=hashlib.new(algorithm,b).hexdigest()
            if actual!=expected:
                raise ValueError("DRYAD SOURCE DIGEST MISMATCH "+filename)
        else:
            raise ValueError("Source file without recognized archive checksum "+filename)
        dest=args.outdir/filename
        dest.write_bytes(b)
        result["file_receipts"].append({
            "filename":filename,"n_bytes":len(b),
            "source_digest_algorithm":algorithm,"source_digest":expected,
            "sha256":hashlib.sha256(b).hexdigest(),
            "url":url,"workbook":inspect_workbook(dest),
        })
    receipt=args.outdir/"original_meta_source_schema_receipt.json"
    receipt.write_text(json.dumps(result,indent=2,sort_keys=True)+"\n",encoding="utf-8")
    print("BEGIN_SCH_META_SOURCE_SCHEMA")
    print(receipt.read_text(encoding="utf-8"))
    print("END_SCH_META_SOURCE_SCHEMA")


if __name__=="__main__":
    main()
